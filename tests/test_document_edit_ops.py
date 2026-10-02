from __future__ import annotations

import hashlib
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from remote_mcp_commander.agent import document_edit_ops
from remote_mcp_commander.agent.document_edit_ops import edit_xlsx_range, replace_docx_text

WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _make_docx(path: Path, *, macro: bool = False) -> None:
    document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{WORD_NS}"><w:body><w:p>
<w:r><w:rPr><w:b/></w:rPr><w:t>Hello </w:t></w:r>
<w:r><w:rPr><w:i/></w:rPr><w:t>World</w:t></w:r>
</w:p></w:body></w:document>'''.encode()
    header = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:hdr xmlns:w="{WORD_NS}"><w:p><w:r><w:t>Header Text</w:t></w:r></w:p></w:hdr>'''.encode()
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", document)
        zf.writestr("word/header1.xml", header)
        if macro:
            zf.writestr("word/vbaProject.bin", b"not-a-real-macro")


def _make_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet["A1"] = "old-a"
    sheet["B1"] = "old-b"
    sheet["A2"] = "styled"
    sheet["A2"].font = Font(bold=True)
    sheet["B2"] = 9
    sheet["C1"] = "=SUM(A1:B1)"
    sheet.merge_cells("D1:E1")
    sheet["D1"] = "merged"
    workbook.save(path)
    workbook.close()


@pytest.mark.asyncio
async def test_docx_replace_handles_split_runs_and_preserves_run_markup(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    _make_docx(path)
    original_hash = _sha256(path)

    result = await replace_docx_text(
        "req",
        str(path),
        "Hello World",
        "Goodbye Team",
        roots=[tmp_path],
        expected_sha256=original_hash,
        expected_replacements=1,
    )

    assert result.rejected is False
    assert result.replacements == 1
    assert result.sha256 == _sha256(path)
    assert result.sha256 != original_hash
    with zipfile.ZipFile(path) as zf:
        xml = zf.read("word/document.xml").decode()
    assert "Goodbye Team" in xml
    assert "Hello " not in xml
    assert "World" not in xml
    assert "<w:b" in xml
    assert "<w:i" in xml
    assert f'xmlns:w="{WORD_NS}"' in xml


@pytest.mark.asyncio
async def test_docx_replace_expected_count_mismatch_is_non_mutating(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    _make_docx(path)
    original_hash = _sha256(path)

    result = await replace_docx_text(
        "req",
        str(path),
        "Hello World",
        "Changed",
        roots=[tmp_path],
        expected_sha256=original_hash,
        expected_replacements=2,
    )

    assert result.rejected is True
    assert "expected 2" in (result.error or "")
    assert _sha256(path) == original_hash
    assert not list(tmp_path.glob(".remote-mcp-docx-*"))


@pytest.mark.asyncio
async def test_docx_edit_rejects_active_content(tmp_path: Path) -> None:
    path = tmp_path / "macro.docx"
    _make_docx(path, macro=True)
    original_hash = _sha256(path)

    result = await replace_docx_text(
        "req",
        str(path),
        "Hello World",
        "Changed",
        roots=[tmp_path],
        expected_sha256=original_hash,
    )

    assert result.rejected is True
    assert "active-content" in (result.error or "")
    assert _sha256(path) == original_hash


@pytest.mark.asyncio
async def test_xlsx_range_edit_preserves_style_and_outside_formula(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    _make_xlsx(path)
    original_hash = _sha256(path)

    result = await edit_xlsx_range(
        "req",
        str(path),
        "Data",
        "A1:B2",
        [[1, 2], [3, 4]],
        roots=[tmp_path],
        expected_sha256=original_hash,
    )

    assert result.rejected is False
    assert result.cells_updated == 4
    assert result.cell_range == "A1:B2"
    assert result.sha256 == _sha256(path)
    workbook = load_workbook(path, data_only=False)
    sheet = workbook["Data"]
    assert [[sheet["A1"].value, sheet["B1"].value], [sheet["A2"].value, sheet["B2"].value]] == [
        [1, 2],
        [3, 4],
    ]
    assert sheet["A2"].font.bold is True
    assert sheet["C1"].value == "=SUM(A1:B1)"
    assert "D1:E1" in {str(item) for item in sheet.merged_cells.ranges}
    workbook.close()


@pytest.mark.asyncio
async def test_xlsx_rejects_formula_injection_without_mutating(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    _make_xlsx(path)
    original_hash = _sha256(path)

    result = await edit_xlsx_range(
        "req",
        str(path),
        "Data",
        "A1:A1",
        [["=WEBSERVICE(\"https://example.invalid\")"]],
        roots=[tmp_path],
        expected_sha256=original_hash,
    )

    assert result.rejected is True
    assert "formula injection" in (result.error or "")
    assert _sha256(path) == original_hash


@pytest.mark.asyncio
async def test_xlsx_rejects_merged_range_and_dimension_mismatch(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    _make_xlsx(path)
    original_hash = _sha256(path)

    merged = await edit_xlsx_range(
        "req",
        str(path),
        "Data",
        "D1:E1",
        [["x", "y"]],
        roots=[tmp_path],
        expected_sha256=original_hash,
    )
    mismatch = await edit_xlsx_range(
        "req2",
        str(path),
        "Data",
        "A1:B2",
        [[1, 2]],
        roots=[tmp_path],
        expected_sha256=original_hash,
    )

    assert merged.rejected is True
    assert "merged cells" in (merged.error or "")
    assert mismatch.rejected is True
    assert "dimensions" in (mismatch.error or "")
    assert _sha256(path) == original_hash


@pytest.mark.asyncio
async def test_document_edits_require_current_sha(tmp_path: Path) -> None:
    docx = tmp_path / "sample.docx"
    xlsx = tmp_path / "sample.xlsx"
    _make_docx(docx)
    _make_xlsx(xlsx)
    wrong_hash = "0" * 64

    docx_result = await replace_docx_text(
        "req",
        str(docx),
        "Hello World",
        "Changed",
        roots=[tmp_path],
        expected_sha256=wrong_hash,
    )
    xlsx_result = await edit_xlsx_range(
        "req2",
        str(xlsx),
        "Data",
        "A1:A1",
        [[1]],
        roots=[tmp_path],
        expected_sha256=wrong_hash,
    )

    assert docx_result.rejected is True
    assert xlsx_result.rejected is True
    assert "changed since preview" in (docx_result.error or "")
    assert "changed since preview" in (xlsx_result.error or "")


@pytest.mark.asyncio
async def test_docx_rejects_control_character_replacement(tmp_path: Path) -> None:
    path = tmp_path / "sample.docx"
    _make_docx(path)
    original_hash = _sha256(path)

    result = await replace_docx_text(
        "req",
        str(path),
        "Hello World",
        "Bad\nText",
        roots=[tmp_path],
        expected_sha256=original_hash,
    )

    assert result.rejected is True
    assert "control characters" in (result.error or "")
    assert _sha256(path) == original_hash


@pytest.mark.asyncio
async def test_xlsx_rejects_non_finite_number(tmp_path: Path) -> None:
    path = tmp_path / "sample.xlsx"
    _make_xlsx(path)
    original_hash = _sha256(path)

    result = await edit_xlsx_range(
        "req",
        str(path),
        "Data",
        "A1:A1",
        [[float("nan")]],
        roots=[tmp_path],
        expected_sha256=original_hash,
    )

    assert result.rejected is True
    assert "non-finite" in (result.error or "")
    assert _sha256(path) == original_hash


@pytest.mark.asyncio
async def test_docx_replace_does_not_cross_structural_tab(tmp_path: Path) -> None:
    path = tmp_path / "tabbed.docx"
    document = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="{WORD_NS}"><w:body><w:p>
<w:r><w:t>Hello</w:t></w:r><w:r><w:tab/></w:r><w:r><w:t>World</w:t></w:r>
</w:p></w:body></w:document>'''.encode()
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", document)
    original_hash = _sha256(path)

    result = await replace_docx_text(
        "req",
        str(path),
        "HelloWorld",
        "Wrong",
        roots=[tmp_path],
        expected_sha256=original_hash,
        expected_replacements=1,
    )

    assert result.rejected is True
    assert "found 0" in (result.error or "")
    assert _sha256(path) == original_hash


@pytest.mark.asyncio
async def test_docx_edit_detects_concurrent_change_before_publish(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "race.docx"
    _make_docx(path)
    original_hash = _sha256(path)
    original_rewrite = document_edit_ops._rewrite_docx_archive

    def racing_rewrite(source: Path, target: Path, replacements: dict[str, bytes]) -> None:
        original_rewrite(source, target, replacements)
        source.write_bytes(b"external-change")

    monkeypatch.setattr(document_edit_ops, "_rewrite_docx_archive", racing_rewrite)
    result = await replace_docx_text(
        "req",
        str(path),
        "Hello World",
        "Changed",
        roots=[tmp_path],
        expected_sha256=original_hash,
    )

    assert result.rejected is True
    assert "changed during edit" in (result.error or "")
    assert path.read_bytes() == b"external-change"
    assert not list(tmp_path.glob(".remote-mcp-docx-*"))


@pytest.mark.asyncio
async def test_xlsx_rejects_dtd_before_openpyxl_parse(tmp_path: Path) -> None:
    path = tmp_path / "dtd.xlsx"
    _make_xlsx(path)
    rewritten = tmp_path / "rewritten.xlsx"
    with zipfile.ZipFile(path, "r") as source, zipfile.ZipFile(
        rewritten, "w", compression=zipfile.ZIP_DEFLATED
    ) as target:
        for info in source.infolist():
            raw = source.read(info)
            if info.filename == "xl/workbook.xml":
                raw = b'<!DOCTYPE workbook [<!ENTITY injected "boom">]>' + raw
            target.writestr(info, raw)
    rewritten.replace(path)
    original_hash = _sha256(path)

    result = await edit_xlsx_range(
        "req",
        str(path),
        "Data",
        "A1:A1",
        [[1]],
        roots=[tmp_path],
        expected_sha256=original_hash,
    )

    assert result.rejected is True
    assert "DTD/entity" in (result.error or "")
    assert _sha256(path) == original_hash
