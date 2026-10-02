from __future__ import annotations

import asyncio
import html
import math
import os
import re
import shutil
import stat
import tempfile
import warnings
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from remote_mcp_commander.agent.document_ops import (
    DOCUMENT_MAX_INPUT_BYTES,
    _parse_cell_range,
    _parse_xml,
    _validate_zip,
)
from remote_mcp_commander.agent.file_ops import file_sha256, resolve_allowed_path, secrets_match
from remote_mcp_commander.protocol import DocumentEditResult, SpreadsheetCellValue

_WORD_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_WORD_PREFIX_RE = re.compile(
    rb"xmlns:([A-Za-z_][A-Za-z0-9_.-]*)=[\"']"
    + re.escape(_WORD_NS.encode("ascii"))
    + rb"[\"']"
)
_DOCX_EDIT_MEMBERS_RE = re.compile(r"^word/(?:document|header\d+|footer\d+)\.xml$")


def _fsync_directory(path: Path) -> None:
    if os.name == "nt":
        return
    fd = os.open(path, os.O_RDONLY | getattr(os, "O_DIRECTORY", 0))
    try:
        os.fsync(fd)
    finally:
        os.close(fd)


def _file_identity(path: Path) -> tuple[int, int, int, int, int, int]:
    info = path.stat()
    return (
        info.st_dev,
        info.st_ino,
        info.st_size,
        info.st_mtime_ns,
        info.st_ctime_ns,
        stat.S_IMODE(info.st_mode),
    )


def _resolve_edit_target(
    raw_path: str,
    roots: list[Path],
    *,
    suffix: str,
    expected_sha256: str,
) -> tuple[Path, int, tuple[int, int, int, int, int, int]]:
    requested = Path(raw_path).expanduser()
    if not requested.is_absolute():
        raise PermissionError("document path must be absolute")
    if requested.is_symlink():
        raise PermissionError("symlink documents are not supported")
    path = resolve_allowed_path(raw_path, roots).resolve(strict=True)
    if not path.is_file():
        raise PermissionError("not a regular file")
    if path.suffix.lower() != suffix:
        raise ValueError(f"expected a {suffix} document")
    if path.stat().st_size > DOCUMENT_MAX_INPUT_BYTES:
        raise ValueError("document exceeds size limit")
    actual_hash = file_sha256(path)
    if not secrets_match(actual_hash, expected_sha256):
        raise ValueError("document changed since preview")
    identity = _file_identity(path)
    return path, stat.S_IMODE(path.stat().st_mode) & 0o777, identity


def _revalidate_before_publish(
    path: Path,
    *,
    expected_sha256: str,
    expected_identity: tuple[int, int, int, int, int, int],
) -> None:
    if path.is_symlink():
        raise PermissionError("document changed to a symlink during edit")
    try:
        current_identity = _file_identity(path)
    except FileNotFoundError as exc:
        raise ValueError("document changed during edit") from exc
    if current_identity != expected_identity:
        raise ValueError("document changed during edit")
    if not secrets_match(file_sha256(path), expected_sha256):
        raise ValueError("document changed during edit")


def _publish_temp(
    path: Path,
    temp_path: Path,
    *,
    mode: int,
    expected_sha256: str,
    expected_identity: tuple[int, int, int, int, int, int],
) -> str:
    if temp_path.stat().st_size > DOCUMENT_MAX_INPUT_BYTES:
        raise ValueError("edited document exceeds size limit")
    with temp_path.open("rb") as handle:
        os.fsync(handle.fileno())
    _revalidate_before_publish(
        path,
        expected_sha256=expected_sha256,
        expected_identity=expected_identity,
    )
    if os.name != "nt":
        os.chmod(temp_path, mode)
    os.replace(temp_path, path)
    _fsync_directory(path.parent)
    return file_sha256(path)


def _reject_macro_content(zf: zipfile.ZipFile) -> None:
    for info in zf.infolist():
        lower = info.filename.lower()
        archived_mode = (info.external_attr >> 16) & 0xFFFF
        if stat.S_ISLNK(archived_mode):
            raise ValueError("symlink Office archive entries are not supported for editing")
        if lower.endswith("vbaproject.bin") or "/activex/" in f"/{lower}":
            raise ValueError("active-content Office documents are not supported for editing")


def _word_prefix(raw: bytes) -> str:
    match = _WORD_PREFIX_RE.search(raw[:32_768])
    if match is None:
        raise ValueError("DOCX WordprocessingML namespace prefix is unavailable")
    return match.group(1).decode("ascii")


def _find_occurrences(text: str, needle: str) -> list[tuple[int, int]]:
    result: list[tuple[int, int]] = []
    offset = 0
    while True:
        index = text.find(needle, offset)
        if index < 0:
            return result
        end = index + len(needle)
        result.append((index, end))
        offset = end


def _replace_paragraph_text(
    paragraph: str, prefix: str, old_text: str, new_text: str
) -> tuple[str, int]:
    text_re = re.compile(
        rf"<{re.escape(prefix)}:t(?P<attrs>[^>]*)>(?P<text>.*?)</{re.escape(prefix)}:t>",
        flags=re.DOTALL,
    )
    matches = list(text_re.finditer(paragraph))
    if not matches:
        return paragraph, 0
    decoded = [html.unescape(match.group("text")) for match in matches]
    structural_gap_re = re.compile(
        rf"<{re.escape(prefix)}:(?:tab|br|cr|drawing|pict|object|sym|"
        rf"footnoteReference|endnoteReference|fldChar|instrText|delText|"
        rf"noBreakHyphen|softHyphen|lastRenderedPageBreak)\b"
    )
    combined_parts: list[str] = []
    boundaries: list[tuple[int, int]] = []
    cursor = 0
    previous_end: int | None = None
    for match, text in zip(matches, decoded, strict=True):
        if previous_end is not None and structural_gap_re.search(
            paragraph[previous_end : match.start()]
        ):
            combined_parts.append("\x00")
            cursor += 1
        boundaries.append((cursor, cursor + len(text)))
        combined_parts.append(text)
        cursor += len(text)
        previous_end = match.end()
    combined = "".join(combined_parts)
    occurrences = _find_occurrences(combined, old_text)
    if not occurrences:
        return paragraph, 0

    operations: list[list[tuple[int, int, str]]] = [[] for _ in matches]
    for start, end in occurrences:
        start_node = next(
            (
                index
                for index, (node_start, node_end) in enumerate(boundaries)
                if node_start <= start < node_end
            ),
            None,
        )
        if start_node is None:
            raise ValueError("DOCX text mapping failed")
        for index, (node_start, node_end) in enumerate(boundaries):
            overlap_start = max(start, node_start)
            overlap_end = min(end, node_end)
            if overlap_start >= overlap_end:
                continue
            operations[index].append(
                (
                    overlap_start - node_start,
                    overlap_end - node_start,
                    new_text if index == start_node else "",
                )
            )

    rewritten_nodes: list[str] = []
    for original, edits in zip(decoded, operations, strict=True):
        if not edits:
            rewritten_nodes.append(original)
            continue
        edits.sort(key=lambda item: item[0])
        pieces: list[str] = []
        cursor = 0
        for start, end, replacement in edits:
            if start < cursor:
                raise ValueError("overlapping DOCX text replacements are not supported")
            pieces.append(original[cursor:start])
            pieces.append(replacement)
            cursor = end
        pieces.append(original[cursor:])
        rewritten_nodes.append("".join(pieces))

    updated = paragraph
    for match, text in reversed(list(zip(matches, rewritten_nodes, strict=True))):
        attrs = match.group("attrs")
        if text and (text[0].isspace() or text[-1].isspace()) and "xml:space=" not in attrs:
            attrs += ' xml:space="preserve"'
        opening = f"<{prefix}:t{attrs}>"
        replacement = opening + html.escape(text, quote=False) + f"</{prefix}:t>"
        updated = updated[: match.start()] + replacement + updated[match.end() :]
    return updated, len(occurrences)


def _replace_docx_member(raw: bytes, old_text: str, new_text: str) -> tuple[bytes, int]:
    _parse_xml(raw)
    prefix = _word_prefix(raw)
    text = raw.decode("utf-8")
    paragraph_re = re.compile(
        rf"<{re.escape(prefix)}:p(?:\s[^>]*)?>.*?</{re.escape(prefix)}:p>",
        flags=re.DOTALL,
    )
    matches = list(paragraph_re.finditer(text))
    updated = text
    replacements = 0
    for match in reversed(matches):
        paragraph, count = _replace_paragraph_text(match.group(0), prefix, old_text, new_text)
        if count:
            replacements += count
            updated = updated[: match.start()] + paragraph + updated[match.end() :]
    encoded = updated.encode("utf-8")
    _parse_xml(encoded)
    return encoded, replacements


def _rewrite_docx_archive(path: Path, temp_path: Path, replacements: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path, "r") as source:
        _validate_zip(source)
        _reject_macro_content(source)
        with zipfile.ZipFile(temp_path, "w") as target:
            for info in source.infolist():
                if info.is_dir():
                    target.writestr(info, b"")
                    continue
                if info.filename in replacements:
                    target.writestr(info, replacements[info.filename])
                    continue
                with source.open(info, "r") as src, target.open(info, "w") as dst:
                    shutil.copyfileobj(src, dst, length=128 * 1024)


def _replace_docx_text_sync(
    request_id: str,
    raw_path: str,
    old_text: str,
    new_text: str,
    *,
    roots: list[Path],
    expected_sha256: str,
    expected_replacements: int,
) -> DocumentEditResult:
    temp_path: Path | None = None
    try:
        if old_text == new_text:
            raise ValueError("DOCX old_text and new_text must differ")
        if any(ord(char) < 0x20 for char in old_text + new_text):
            raise ValueError("DOCX edit text must not contain control characters")
        path, mode, identity = _resolve_edit_target(
            raw_path, roots, suffix=".docx", expected_sha256=expected_sha256
        )
        member_updates: dict[str, bytes] = {}
        total = 0
        with zipfile.ZipFile(path, "r") as zf:
            _validate_zip(zf)
            _reject_macro_content(zf)
            for info in zf.infolist():
                if not _DOCX_EDIT_MEMBERS_RE.fullmatch(info.filename):
                    continue
                raw = zf.read(info)
                updated, count = _replace_docx_member(raw, old_text, new_text)
                if count:
                    member_updates[info.filename] = updated
                    total += count
        if total != expected_replacements:
            raise ValueError(
                f"expected {expected_replacements} DOCX replacements but found {total}"
            )

        fd, temp_name = tempfile.mkstemp(
            prefix=".remote-mcp-docx-", suffix=".docx", dir=path.parent
        )
        os.close(fd)
        temp_path = Path(temp_name)
        _rewrite_docx_archive(path, temp_path, member_updates)
        sha256 = _publish_temp(
            path,
            temp_path,
            mode=mode,
            expected_sha256=expected_sha256,
            expected_identity=identity,
        )
        temp_path = None
        return DocumentEditResult(
            request_id=request_id,
            path=str(path),
            kind="docx",
            replacements=total,
            bytes_written=path.stat().st_size,
            sha256=sha256,
        )
    except (OSError, PermissionError, UnicodeError, ValueError, zipfile.BadZipFile) as exc:
        return DocumentEditResult(
            request_id=request_id,
            path=raw_path,
            kind="docx",
            rejected=True,
            error=str(exc),
        )
    finally:
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)


def _validate_xlsx_values(values: list[list[SpreadsheetCellValue]]) -> tuple[int, int]:
    if not values or not values[0]:
        raise ValueError("XLSX values must contain at least one cell")
    width = len(values[0])
    for row in values:
        if len(row) != width:
            raise ValueError("XLSX values must be a rectangular 2D array")
        for value in row:
            if isinstance(value, str) and value.startswith("="):
                raise ValueError("XLSX formula injection is not supported by this editor")
            if isinstance(value, float) and not math.isfinite(value):
                raise ValueError("XLSX non-finite numeric values are not supported")
    return len(values), width


def _edit_xlsx_range_sync(
    request_id: str,
    raw_path: str,
    sheet: str,
    cell_range: str,
    values: list[list[SpreadsheetCellValue]],
    *,
    roots: list[Path],
    expected_sha256: str,
) -> DocumentEditResult:
    temp_path: Path | None = None
    workbook = None
    try:
        path, mode, identity = _resolve_edit_target(
            raw_path, roots, suffix=".xlsx", expected_sha256=expected_sha256
        )
        with zipfile.ZipFile(path, "r") as zf:
            _validate_zip(zf)
            _reject_macro_content(zf)
        rows, columns = _validate_xlsx_values(values)
        start_col, start_row, end_col, end_row, normalized = _parse_cell_range(cell_range)
        if end_row - start_row + 1 != rows or end_col - start_col + 1 != columns:
            raise ValueError("XLSX range dimensions must exactly match the 2D values array")

        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            workbook = load_workbook(path, read_only=False, data_only=False, keep_links=True)
            if sheet not in workbook.sheetnames:
                raise ValueError(f"XLSX worksheet not found: {sheet}")
            worksheet = workbook[sheet]
            for merged in worksheet.merged_cells.ranges:
                if not (
                    merged.max_col < start_col + 1
                    or merged.min_col > end_col + 1
                    or merged.max_row < start_row
                    or merged.min_row > end_row
                ):
                    raise ValueError("XLSX edit range intersects merged cells")
            for row_offset, row_values in enumerate(values):
                for column_offset, value in enumerate(row_values):
                    worksheet.cell(
                        row=start_row + row_offset,
                        column=start_col + column_offset + 1,
                    ).value = value

            fd, temp_name = tempfile.mkstemp(
                prefix=".remote-mcp-xlsx-", suffix=".xlsx", dir=path.parent
            )
            os.close(fd)
            temp_path = Path(temp_name)
            workbook.save(temp_path)
            risky = [
                str(item.message)
                for item in caught
                if "not supported" in str(item.message).lower()
                or "will be removed" in str(item.message).lower()
            ]
            if risky:
                raise ValueError(f"XLSX contains unsupported features: {risky[0][:256]}")
        workbook.close()
        workbook = None

        sha256 = _publish_temp(
            path,
            temp_path,
            mode=mode,
            expected_sha256=expected_sha256,
            expected_identity=identity,
        )
        temp_path = None
        return DocumentEditResult(
            request_id=request_id,
            path=str(path),
            kind="xlsx",
            sheet=sheet,
            cell_range=normalized,
            cells_updated=rows * columns,
            bytes_written=path.stat().st_size,
            sha256=sha256,
        )
    except (
        InvalidFileException,
        KeyError,
        OSError,
        OverflowError,
        PermissionError,
        TypeError,
        ValueError,
        zipfile.BadZipFile,
    ) as exc:
        return DocumentEditResult(
            request_id=request_id,
            path=raw_path,
            kind="xlsx",
            rejected=True,
            error=str(exc),
        )
    finally:
        if workbook is not None:
            workbook.close()
        if temp_path is not None:
            temp_path.unlink(missing_ok=True)


async def replace_docx_text(
    request_id: str,
    raw_path: str,
    old_text: str,
    new_text: str,
    *,
    roots: list[Path],
    expected_sha256: str,
    expected_replacements: int = 1,
) -> DocumentEditResult:
    return await asyncio.to_thread(
        _replace_docx_text_sync,
        request_id,
        raw_path,
        old_text,
        new_text,
        roots=roots,
        expected_sha256=expected_sha256,
        expected_replacements=expected_replacements,
    )


async def edit_xlsx_range(
    request_id: str,
    raw_path: str,
    sheet: str,
    cell_range: str,
    values: list[list[SpreadsheetCellValue]],
    *,
    roots: list[Path],
    expected_sha256: str,
) -> DocumentEditResult:
    return await asyncio.to_thread(
        _edit_xlsx_range_sync,
        request_id,
        raw_path,
        sheet,
        cell_range,
        values,
        roots=roots,
        expected_sha256=expected_sha256,
    )
